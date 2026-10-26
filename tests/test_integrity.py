"""행 수·금액 대사 — 숫자가 안 맞으면 업로드본을 쓰지 않는다. 합성 데이터만 쓴다."""
from decimal import Decimal

import openpyxl
import pytest

from kafa.integrity import (HOLD_SUFFIX, Count, IntegrityError, check_store,
                            check_upload_files, hold, read_back, render, tally_batch)
from kafa.io_wehago.reader import DownloadRows, read_download_xlsx
from kafa.io_wehago.schema import INPUT_COLUMNS
from kafa.io_wehago.writer import to_output_row, write_upload_xls
from kafa.rules.engine import classify_row
from kafa.rules.models import InputRow


def _in(거래처="행복상사", 합계=5500, 전표상태="", 일자="03-15"):
    return InputRow(연도="2026", 일자=일자, 거래처=거래처, 공급가액=Decimal(합계),
                    세액=Decimal(0), 합계=Decimal(합계), 국세청="공제", 유형="카과",
                    차변계정="(판)복리후생비", 전표상태=전표상태,
                    사업자등록번호="111-11-11119")


def _run(rows):
    classified = [classify_row(r, client_type="corporate") for r in rows]
    out = [to_output_row(c) for c in classified if not c.skipped]
    return classified, out


# ── 셈 ──

def test_count_compares_amounts_to_the_won():
    assert Count(1, Decimal("7900")).same(Count(1, Decimal("7900.0")))
    assert not Count(1, Decimal("7900")).same(Count(1, Decimal("7901")))
    assert not Count(2, Decimal("7900")).same(Count(1, Decimal("7900")))
    assert str(Count(3, Decimal("12000"))) == "3건 12,000원"


# ── 1~3: 읽은 행 → 분류 → 업로드 + 스킵 ──

def test_a_clean_batch_balances():
    rows = [_in(), _in("다른가게", 9900), _in("중복가게", 1100, "중복전표")]
    classified, out = _run(rows)
    t = tally_batch(rows, classified, out)
    assert t.ok
    assert t.입력 == Count(3, Decimal(16500))
    assert t.업로드.건수 == 2 and t.스킵.건수 == 1
    assert t.스킵_사유["중복전표"] == Count(1, Decimal(1100))
    assert "대사 OK" in t.line()


def test_a_row_lost_between_reading_and_classifying_is_caught():
    rows = [_in(), _in("다른가게", 9900)]
    classified, out = _run(rows)
    t = tally_batch(rows, classified[:1], out[:1])
    assert not t.ok and any("분류된 행" in p for p in t.problems)


def test_a_row_lost_before_the_upload_is_caught():
    rows = [_in(), _in("다른가게", 9900)]
    classified, out = _run(rows)
    t = tally_batch(rows, classified, out[:1])
    assert not t.ok and any("업로드+스킵" in p for p in t.problems)


def test_an_amount_changed_on_the_way_is_caught():
    """건수는 맞아도 금액이 바뀌면 잡는다."""
    rows = [_in()]
    classified, out = _run(rows)
    out[0].합계 = Decimal(5000)
    assert not tally_batch(rows, classified, out).ok


def test_rows_set_aside_on_purpose_are_counted():
    """보충본의 --skip-unresolved 처럼 일부러 뺀 행은 '따로 뺀 행'으로 맞춘다."""
    rows = [_in(), _in("다른가게", 9900)]
    classified, out = _run(rows)
    t = tally_batch(rows, classified, out[:1], held=classified[1:])
    assert t.ok and t.보류 == Count(1, Decimal(9900))


# ── 파일이 스스로 밝힌 총계 ──

def _download(tmp_path, n_rows, label, total):
    p = tmp_path / "d.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(INPUT_COLUMNS)
    for i in range(n_rows):
        r = ["2026", f"01-0{i + 1}", "A", f"가맹점{i}", "법인", "", 1000, 100, 0, 1100,
             "공제", "도매", "문구", "카과", "(판)소모품비", "", "", "", "111-11-11119"]
        ws.append(r)
    for text, amount in [("카드사별  매입 : 1건", "1,100"), (label, total)]:
        s = [""] * len(INPUT_COLUMNS); s[1] = text; s[9] = amount
        ws.append(s)
    wb.save(p)
    return p


def test_reader_picks_up_the_file_total_not_a_card_subtotal(tmp_path):
    rows = read_download_xlsx(_download(tmp_path, 2, "합계(카드사:1) : 2건", "2,200"))
    assert isinstance(rows, DownloadRows) and len(rows) == 2
    assert rows.declared_count == 2 and rows.declared_amount == Decimal(2200)


def test_file_total_count_mismatch_stops_the_run(tmp_path):
    rows = read_download_xlsx(_download(tmp_path, 2, "합계(카드사:1) : 3건", "3,300"))
    classified, out = _run(rows)
    t = tally_batch(rows, classified, out)
    assert not t.ok and any("하단 총계 3건" in p for p in t.problems)


def test_file_total_amount_mismatch_only_warns_by_default(tmp_path):
    """금액 의미는 실샘플로 아직 확인 전이라 경고만 한다(config 로 fail 전환)."""
    rows = read_download_xlsx(_download(tmp_path, 2, "합계(카드사:1) : 2건", "9,999"))
    classified, out = _run(rows)
    t = tally_batch(rows, classified, out)
    assert t.ok and t.warnings and "경고" in t.line()


def test_a_file_without_a_total_line_is_not_judged_on_it():
    rows = [_in()]
    classified, out = _run(rows)
    t = tally_batch(rows, classified, out)
    assert t.파일_총계 is None and t.ok


# ── 4: 업로드 파일 다시 읽기 ──

def test_read_back_counts_every_split_part(tmp_path):
    rows = [_in(f"가게{i}", 1000 + i) for i in range(30)]
    classified, out = _run(rows)
    files = write_upload_xls(out, tmp_path / "u.xls", strict=False, max_bytes=8000)
    assert len(files) > 1                                   # 분할됨
    assert read_back(files) == Count(30, sum(Decimal(1000 + i) for i in range(30)))


def test_an_upload_file_that_lost_rows_is_caught(tmp_path):
    rows = [_in(), _in("다른가게", 9900)]
    classified, out = _run(rows)
    t = tally_batch(rows, classified, out)
    files = write_upload_xls(out[:1], tmp_path / "u.xls", strict=False)   # 한 줄 빠진 파일
    check_upload_files(t, files)
    assert not t.ok and any("다시 읽으니" in p for p in t.problems)


def test_a_broken_upload_is_held_so_it_cannot_be_uploaded(tmp_path):
    f = tmp_path / "u.xls"
    f.write_bytes(b"x")
    held = hold([f])
    assert not f.exists() and held[0].name == "u.xls" + HOLD_SUFFIX


# ── 5: DB 적재 ──

def test_keys_that_collide_are_caught_with_the_lost_count():
    rows = [_in(), _in(), _in()]
    classified, out = _run(rows)
    t = tally_batch(rows, classified, out)
    check_store(t, ["k1", "k1", "k1"], present=1)        # 고치기 전의 voucher_key 처럼
    assert not t.ok and any("2건이 같은 키로 겹침" in p for p in t.problems)


def test_keys_missing_from_the_store_are_caught():
    rows = [_in(), _in("다른가게", 9900)]
    classified, out = _run(rows)
    t = tally_batch(rows, classified, out)
    check_store(t, ["k1", "k2"], present=1)
    assert not t.ok and any("실제 적재 1개" in p for p in t.problems)


# ── 보고서에는 실명이 없다 ──

def test_the_report_carries_counts_not_names():
    rows = [_in("행복상사", 5500, "중복전표")]
    classified, out = _run(rows)
    text = render(tally_batch(rows, classified, out))
    assert "행복상사" not in text and "111-11-11119" not in text
    assert "중복전표" in text and "1건 5,500원" in text


# ── 처리 경로에 끼워져 있는가 ──

def test_process_rows_refuses_to_write_an_upload_that_does_not_balance(tmp_path):
    from kafa.cli import process_rows

    rows = DownloadRows([_in(), _in("다른가게", 9900)])
    rows.declared_count = 5                                 # 파일은 5건이라는데 2건만 읽힘
    out = tmp_path / "x_upload.xls"
    with pytest.raises(IntegrityError) as e:
        process_rows(rows, out)
    assert not out.exists()                                 # 업로드본을 아예 쓰지 않는다
    assert (tmp_path / "x_upload_대사.txt").exists()
    assert "하단 총계 5건" in str(e.value)


def test_process_rows_reports_a_clean_tally(tmp_path):
    from kafa.cli import process_rows

    res = process_rows([_in(), _in("다른가게", 9900)], tmp_path / "x_upload.xls")
    assert res["tally"].ok and res["tally"].업로드_파일 == Count(2, Decimal(15400))


# ── 완료 기준: 2026-09 의 169건 유실을 재현하면 실패로 잡혀야 한다 ──

def _inbox(tmp_path, n_same):
    inbox = tmp_path / "inbox"
    p = inbox / "고객001" / "3월.xlsx"
    p.parent.mkdir(parents=True)
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(INPUT_COLUMNS)
    for _ in range(n_same):           # 같은 날 같은 가게 같은 금액 — 통행료·편의점처럼
        ws.append(["2026", "03-15", "A", "한국도로공사", "법인", "", 818, 82, 0, 900,
                   "공제", "서비스", "유료도로운영", "카과", "(판)여비교통비", "", "",
                   "", "111-11-11119"])
    wb.save(p)
    return inbox, p


def test_the_september_key_collision_is_now_caught(tmp_path, monkeypatch):
    """고치기 전 voucher_key(순번 없음)로 돌리면 파이프라인이 그 파일을 실패로 멈춘다."""
    from kafa.pipeline.runner import run_pipeline
    from kafa.store import db as store_db

    original = store_db._voucher_key
    monkeypatch.setattr(store_db, "_voucher_key",
                        lambda cid, c, occurrence=0: original(cid, c, 0))
    inbox, src = _inbox(tmp_path, 3)
    out = tmp_path / "out"

    res = run_pipeline(inbox, out)

    assert not res.ok and len(res.failures) == 1
    message = next(iter(res.failures.values()))
    assert "IntegrityError" in message and "2건이 같은 키로 겹침" in message
    base = out / "고객001" / "2026-03"
    assert not (base / "3월_upload.xls").exists()                   # 올릴 수 있는 파일 없음
    assert (base / ("3월_upload.xls" + HOLD_SUFFIX)).exists()       # 보류로 남는다
    assert (base / "3월_upload_대사.txt").exists()
    assert src.exists()                         # 원본은 보관함으로 안 옮긴다 → 고친 뒤 재실행


def test_with_the_fix_the_same_file_balances(tmp_path):
    from kafa.pipeline.runner import run_pipeline

    inbox, src = _inbox(tmp_path, 3)
    res = run_pipeline(inbox, tmp_path / "out")
    assert res.ok and res.outcomes[0].inserted == 3
    assert res.outcomes[0].integrity.startswith("대사 OK — 입력 3 = 업로드 3 + 스킵 0")


# ── 2차 중복 가드: 같은 파일 안의 반복 결제는 별개 거래, 같은 파일 재처리는 막는다 ──

def test_repeat_purchases_in_one_file_all_reach_the_upload(tmp_path):
    """2026-09 의 169건은 이 가드가 두 번째 행부터 '중복'으로 지워서 업로드본에서 빠졌다."""
    from kafa.cli import classify_rows
    from kafa.dup_guard import DupGuard

    dup = DupGuard(tmp_path / "dup.json")
    classified, skipped = classify_rows([_in("한국도로공사", 900)] * 3, dup=dup,
                                        client_type="corporate")
    assert skipped == 0 and all(not c.skipped for c in classified)


def test_reprocessing_the_same_file_is_still_blocked(tmp_path):
    from kafa.cli import classify_rows
    from kafa.dup_guard import DupGuard

    rows = [_in("한국도로공사", 900)] * 3
    first = DupGuard(tmp_path / "dup.json")
    classify_rows(rows, dup=first, client_type="corporate")
    first.flush()                                  # 첫 실행이 끝까지 성공했다
    again, skipped = classify_rows(rows, dup=DupGuard(tmp_path / "dup.json"),
                                   client_type="corporate")
    assert skipped == 3 and {c.skip_reason for c in again} == {"dup_guard(2차)"}


def test_the_first_occurrence_keeps_its_old_key():
    """기존 dup.json 에 남은 키와 맞도록 첫 행의 키 모양은 그대로다."""
    from kafa.dup_guard import make_key

    args = ("2026-03-15", "가게", "111-11-11119", Decimal(900), "")
    assert make_key(*args) == make_key(*args, occurrence=0)
    assert make_key(*args, occurrence=1) != make_key(*args)


def test_the_report_says_whether_the_upload_was_held_or_never_made():
    rows = [_in()]
    classified, out = _run(rows)
    t = tally_batch(rows, classified, out[:0])
    assert "만들지 않았습니다" in render(t) and HOLD_SUFFIX not in render(t)
    assert "보류했습니다" in render(t, held=True) and HOLD_SUFFIX in render(t, held=True)


def test_a_failed_run_does_not_poison_the_rerun(tmp_path):
    """대사가 깨진 실행이 중복 기록을 남기면, 고친 뒤 다시 돌릴 때 전 행이 '중복'으로 빠져
    업로드 0 + 스킵 N 으로 대사가 'OK' 가 된다. 실패한 실행은 기록을 남기지 않아야 한다."""
    from kafa.pipeline.runner import run_pipeline

    inbox, src = _inbox(tmp_path, 2)
    wb = openpyxl.load_workbook(src); ws = wb.active
    s = [""] * len(INPUT_COLUMNS); s[1] = "합계(카드사:1) : 3건"; s[9] = "2,700"
    ws.append(s); wb.save(src)                     # 하단 총계가 안 맞는 파일
    out = tmp_path / "out"
    assert not run_pipeline(inbox, out).ok

    ws.delete_rows(ws.max_row); wb.save(src)       # 원인을 고쳤다
    again = run_pipeline(inbox, out)
    assert again.ok
    o = again.outcomes[0]
    assert o.written == 2 and o.skipped == 0       # 빈 업로드본이 통과하지 않는다
