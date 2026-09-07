"""수임처에게 보낼 자료요청서 — 앱 주문내역은 본인만 볼 수 있다.

쿠팡·네이버페이·카카오페이 건은 실판매자와 품목이 **그 앱의 주문내역**에 남는다.
데이터 질은 제일 좋은데 수임처 대표 개인 계정이라 세무대리인이 볼 수 없다.
그래서 조회하는 대신 **물어본다** — 날짜·금액을 적어 주고 채워 달라고 한다.

승인번호가 필요 없다. 본인은 자기 주문목록에서 날짜·금액으로 바로 찾는다.
그래서 1단계(카드사 이용내역)를 기다릴 것 없이 지금 바로 보낼 수 있다.

**수임처마다 파일을 따로 낸다.** 한 파일에 여러 수임처를 담으면 남의 거래내역이
보인다 — 이건 편의 문제가 아니라 지켜야 할 선이다.
"""
from __future__ import annotations

import csv
import re
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from kafa.config_loader import load_pg_sites
from kafa.lookup.merge import FOUND
from kafa.lookup.statements import normalize_amount
from kafa.lookup.stage2 import APP_ORDER, route_for

_UNSAFE = re.compile(r'[\\/:*?"<>|]+')


@dataclass
class Ask:
    """수임처에게 물어볼 한 건."""
    거래일자: str = ""
    금액: Decimal = Decimal(0)
    창구: str = ""
    안내: str = ""
    대행사: str = ""

    @property
    def 구분(self) -> str:
        return "환불" if self.금액 < 0 else "결제"


@dataclass
class ClientAsk:
    """수임처 한 곳에 보낼 요청서."""
    번호: str = ""
    수임처: str = ""
    항목: list[Ask] = field(default_factory=list)
    파일: Path | None = None

    @property
    def 건수(self) -> int:
        return len(self.항목)

    @property
    def 금액합(self) -> Decimal:
        return sum((abs(a.금액) for a in self.항목), Decimal(0))

    @property
    def 기간(self) -> str:
        days = sorted(a.거래일자 for a in self.항목 if a.거래일자)
        return f"{days[0]} ~ {days[-1]}" if days else ""


def _safe(name: str) -> str:
    return _UNSAFE.sub("_", (name or "수임처").strip()) or "수임처"


def collect(source_csv: str | Path, *, config_dir: str | None = None) -> list[ClientAsk]:
    """의뢰 목록(또는 1단계 결과)에서 '수임처에게 물어야 하는 건'만 모은다."""
    spec = (load_pg_sites(config_dir) or {}).get("routes", {}).get(APP_ORDER, {})
    안내표 = spec.get("안내") or {}

    골라낸: list[tuple[str, str, str, Ask]] = []
    with Path(source_csv).open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            if row.get("결과") == FOUND:
                continue                      # 1단계에서 이미 풀린 건은 묻지 않는다
            대행사 = row.get("대행사") or row.get("거래처") or ""
            경로, site = route_for(대행사, config_dir=config_dir)
            if 경로 != APP_ORDER:
                continue
            창구 = site.get("창구", "")
            key = row.get("client_id") or row.get("수임처") or row.get("번호") or ""
            골라낸.append((key, row.get("번호", ""), row.get("수임처", ""), Ask(
                거래일자=row.get("거래일자", ""),
                금액=normalize_amount(row.get("합계", "")),
                창구=창구, 안내=str(안내표.get(창구, "")), 대행사=대행사)))

    # 결제와 취소가 짝을 이뤄 순액이 0인 건은 묻지 않는다 — 물어봐야 살 게 없다.
    def 짝(key: str, a: Ask) -> tuple:
        return (key, a.거래일자, a.대행사, abs(a.금액))

    양수 = {짝(k, a) for k, _, _, a in 골라낸 if a.금액 > 0}
    음수 = {짝(k, a) for k, _, _, a in 골라낸 if a.금액 < 0}
    상쇄 = 양수 & 음수

    per: dict[str, ClientAsk] = {}
    for key, 번호, 수임처, a in 골라낸:
        if a.금액 == 0 or 짝(key, a) in 상쇄:
            continue
        entry = per.get(key)
        if entry is None:
            entry = per[key] = ClientAsk(번호=번호, 수임처=수임처)
        entry.항목.append(a)

    out = sorted(per.values(), key=lambda c: -c.건수)
    for c in out:
        c.항목.sort(key=lambda a: (a.창구, a.거래일자))
    return out


def _write_xlsx(client: ClientAsk, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "자료요청"

    제목 = ws.cell(row=1, column=1, value=f"{client.수임처} — 카드 결제 내역 확인 요청")
    제목.font = Font(size=14, bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)

    안내문 = [
        f"기간 {client.기간} / {client.건수}건 / 합계 {int(client.금액합):,}원",
        "",
        "아래 결제 건은 카드 내역에 결제대행사 이름만 남아 실제로 어디서 무엇을 사셨는지",
        "저희 쪽에서 확인할 수 없습니다. 각 앱의 주문내역에서 찾아 오른쪽 두 칸만 채워 주세요.",
        "정확한 계정 처리를 위해 필요하며, 금액이 큰 건부터 채우셔도 됩니다.",
        "(구분이 '환불'인 줄은 결제 취소 건입니다. 무엇을 취소하셨는지 적어 주세요.)",
    ]
    for i, line in enumerate(안내문, start=2):
        cell = ws.cell(row=i, column=1, value=line)
        if i == 2:
            cell.font = Font(bold=True)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=7)

    head_row = len(안내문) + 3
    headers = ["결제일", "구분", "금액", "결제수단", "어디서 찾나",
               "실제 구입처", "무엇을 구입"]
    fill = PatternFill("solid", fgColor="E8EDF3")
    thin = Side(style="thin", color="B9C4D0")
    for col, name in enumerate(headers, start=1):
        c = ws.cell(row=head_row, column=col, value=name)
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
        c.border = Border(bottom=thin, top=thin, left=thin, right=thin)

    답칸 = PatternFill("solid", fgColor="FFF9E6")
    for offset, item in enumerate(client.항목, start=1):
        r = head_row + offset
        ws.cell(row=r, column=1, value=item.거래일자)
        구분 = ws.cell(row=r, column=2, value=item.구분)
        if item.구분 == "환불":
            구분.font = Font(color="B03A2E", bold=True)
        금액 = ws.cell(row=r, column=3, value=int(item.금액))
        금액.number_format = "#,##0;[Red]-#,##0"
        ws.cell(row=r, column=4, value=item.창구)
        ws.cell(row=r, column=5, value=item.안내)
        for col in (6, 7):
            blank = ws.cell(row=r, column=col, value="")
            blank.fill = 답칸
            blank.border = Border(left=thin, right=thin, bottom=thin)

    for col, width in enumerate([12, 7, 12, 16, 46, 26, 30], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_requests(source_csv: str | Path, out_dir: str | Path, *,
                   config_dir: str | None = None) -> list[ClientAsk]:
    """수임처마다 자료요청서 한 부씩 + 담당자용 발송 목록."""
    clients = collect(source_csv, config_dir=config_dir)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)

    for i, c in enumerate(clients, start=1):
        번호 = (c.번호 or str(i)).strip() or str(i)
        name = f"자료요청_{번호.zfill(3)}_{_safe(c.수임처)}.xlsx"
        c.파일 = out / name
        _write_xlsx(c, c.파일)

    with (out / "발송목록.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["번호", "수임처", "건수", "금액합", "기간", "파일"])
        for c in clients:
            w.writerow([c.번호, c.수임처, c.건수, int(c.금액합), c.기간,
                        c.파일.name if c.파일 else ""])
    return clients
