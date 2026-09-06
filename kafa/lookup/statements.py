"""카드사 이용내역(명세서) 읽기 — 카드사마다 모양이 달라 유연하게 읽는다.

헤더 문구도, 헤더가 몇 번째 줄에 있는지도 카드사마다 다르다(제목·조회기간 안내가
위에 붙는 경우가 흔하다). 그래서 **날짜 칸과 금액 칸이 같이 있는 줄**을 헤더로 본다.
문구 목록은 `config/lookup/statements.yaml` 에 있다 — 코드가 아니라 거기서 늘린다.
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from kafa.config_loader import load_lookup_spec

_DATE = re.compile(r"(\d{4})[.\-/년\s]*(\d{1,2})[.\-/월\s]*(\d{1,2})")
_NON_NUM = re.compile(r"[^\d.\-]")


@dataclass
class StatementRow:
    """이용내역 한 줄에서 우리가 쓰는 것만."""
    날짜: str = ""
    금액: Decimal = Decimal(0)
    가맹점: str = ""
    승인번호: str = ""
    카드: str = ""
    사업자번호: str = ""
    출처: str = ""


def _spec(config_dir: str | None) -> dict:
    return load_lookup_spec(config_dir) or {}


def normalize_date(text: object) -> str:
    """어떤 표기로 오든 YYYY-MM-DD 로. 못 읽으면 빈 문자열."""
    raw = str(text or "").strip()
    if not raw:
        return ""
    digits = re.sub(r"\D", "", raw)
    if len(digits) >= 8 and not _DATE.search(raw):
        y, m, d = digits[:4], digits[4:6], digits[6:8]
    else:
        m_ = _DATE.search(raw)
        if not m_:
            return ""
        y, m, d = m_.group(1), m_.group(2), m_.group(3)
    try:
        if not (1900 <= int(y) <= 2200 and 1 <= int(m) <= 12 and 1 <= int(d) <= 31):
            return ""
    except ValueError:
        return ""
    return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"


def normalize_amount(text: object) -> Decimal:
    """'1,234원', '(1,234)', '-1234' 을 Decimal 로. 괄호는 음수로 본다."""
    raw = str(text or "").strip()
    if not raw:
        return Decimal(0)
    negative = raw.startswith("(") and raw.endswith(")")
    cleaned = _NON_NUM.sub("", raw)
    if not cleaned or cleaned in {"-", ".", "-."}:
        return Decimal(0)
    try:
        value = Decimal(cleaned)
    except InvalidOperation:
        return Decimal(0)
    return -value if negative and value > 0 else value


def _alias_map(config_dir: str | None) -> dict[str, str]:
    """헤더 문구 → 우리 필드 이름. 공백·괄호를 지우고 비교한다."""
    out: dict[str, str] = {}
    for field, names in (_spec(config_dir).get("columns") or {}).items():
        for name in names or []:
            out[re.sub(r"[\s()（）_\-]", "", str(name))] = field
    return out


def _cell(value: object) -> str:
    return "" if value is None else str(value).strip()


def find_header(table: list[list], config_dir: str | None = None):
    """(헤더 줄 번호, {필드: 열번호}) — 날짜와 금액이 같이 있는 첫 줄이 헤더다."""
    aliases = _alias_map(config_dir)
    limit = int(_spec(config_dir).get("header_scan_rows", 12))
    for index, row in enumerate(table[:limit]):
        mapping: dict[str, int] = {}
        for col, cell in enumerate(row):
            key = re.sub(r"[\s()（）_\-]", "", _cell(cell))
            field = aliases.get(key)
            if field and field not in mapping:
                mapping[field] = col
        if "날짜" in mapping and "금액" in mapping:
            return index, mapping
    return -1, {}


def _rows_of(path: Path) -> list[list]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            return [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
        finally:
            wb.close()
    if suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(str(path))
        sheet = book.sheet_by_index(0)
        return [sheet.row_values(i) for i in range(sheet.nrows)]
    for encoding in ("utf-8-sig", "cp949", "utf-8"):
        try:
            with path.open(encoding=encoding, newline="") as fh:
                return [row for row in csv.reader(fh)]
        except UnicodeDecodeError:
            continue
    raise ValueError(f"인코딩을 알 수 없습니다: {path.name}")


def read_statement(path: str | Path, *, config_dir: str | None = None) -> list[StatementRow]:
    """이용내역 파일 하나를 StatementRow 목록으로. 헤더를 못 찾으면 ValueError."""
    file = Path(path)
    table = _rows_of(file)
    index, mapping = find_header(table, config_dir)
    if index < 0:
        raise ValueError(
            f"{file.name}: 날짜·금액 컬럼을 못 찾았습니다 — "
            "config/lookup/statements.yaml 에 그 카드사의 헤더 문구를 추가하세요")

    out: list[StatementRow] = []
    for row in table[index + 1:]:
        def get(field: str) -> str:
            col = mapping.get(field)
            return _cell(row[col]) if col is not None and col < len(row) else ""

        날짜 = normalize_date(get("날짜"))
        금액 = normalize_amount(get("금액"))
        if not 날짜 or 금액 == 0:
            continue
        out.append(StatementRow(
            날짜=날짜, 금액=금액, 가맹점=get("가맹점"), 승인번호=get("승인번호"),
            카드=get("카드"), 사업자번호=get("사업자번호"), 출처=file.name))
    return out
